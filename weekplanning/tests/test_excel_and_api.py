from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.db import Database, reset_db_singleton
from app.excel_import import import_excel, seed_if_needed
from app.main import app
from app.pdf_extract import extract_from_page_texts
from app.services import save_extracted_project, upsert_profile
from tests.test_pdf_extract import GOLD_A50, _page, _write_pdf


ROOT = Path(__file__).resolve().parents[1]
SEED = ROOT / "seed" / "Inrichting_Elace.xlsx"


def test_seed_inrichting_excel(tmp_path):
    db = Database(tmp_path / "t.db")
    result = seed_if_needed(SEED, db=db)
    assert result["seeded"] is True
    assert result["layout"] == "inrichting"
    numbers = [r["number"] for r in db.fetchall("SELECT number FROM profiles ORDER BY number")]
    assert numbers == ["AA4410", "AA5110", "EB", "RT"]
    rt = db.fetchone("SELECT comet_followup FROM profiles WHERE number = 'RT'")
    assert rt["comet_followup"] == "X"
    again = seed_if_needed(SEED, db=db)
    assert again["seeded"] is False
    assert db.fetchone("SELECT COUNT(*) AS c FROM profiles")["c"] == 4


def test_zaagmatrix_excel(tmp_path):
    path = tmp_path / "matrix.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["B1"] = "Quadra"
    ws["D1"] = "BJM"
    ws["F1"] = "Schirmer"
    ws["H1"] = "Comet"
    ws["B2"] = "Beperkt"
    ws["C2"] = "Niet"
    ws["D2"] = "Beperkt"
    ws["E2"] = "Niet"
    ws["F2"] = "Beperkt"
    ws["G2"] = "Niet"
    ws["A3"] = "0115850"
    ws["C3"] = "x"  # niet op Quadra
    ws["D3"] = "x"  # beperkt BJM
    ws["H3"] = "Q"
    wb.save(path)
    db = Database(tmp_path / "t.db")
    result = import_excel(path, db=db)
    db.commit()
    assert result["layout"] == "zaagmatrix"
    p = db.fetchone("SELECT * FROM profiles WHERE number = '0115850'")
    assert p["comet_followup"] == "Q"
    flags = {
        r["code"]: (r["beperkt"], r["niet"])
        for r in db.fetchall(
            """
            SELECT m.code, pm.beperkt, pm.niet
            FROM profile_machine pm JOIN machines m ON m.id = pm.machine_id
            WHERE pm.profile_id = ?
            """,
            (p["id"],),
        )
    }
    assert flags["quadra"] == (0, 1)
    assert flags["bjm"] == (1, 0)
    assert flags["schirmer"] == (0, 0)


def test_project_crud_and_unknown_flag(tmp_path):
    db = Database(tmp_path / "t.db")
    upsert_profile(db, "0115850", "", "Q", {"quadra": {"niet": True}})
    extracted = extract_from_page_texts([_page("012510375A50", GOLD_A50[:3])])
    project = save_extracted_project(db, extracted, "demo.pdf")
    assert project["number"] == "012510375A50"
    assert project["lines"][0]["known"] is True
    assert project["lines"][0]["machines"][0]["status"] == "niet"
    assert project["lines"][1]["known"] is False
    from app.services import add_line, delete_line, update_line_qty
    project = update_line_qty(db, project["lines"][0]["id"], 9)
    assert project["lines"][0]["aantal"] == 9
    project = add_line(db, project["id"], "9999999", 3)
    assert project["lines"][-1]["profile_number"] == "9999999"
    last_id = project["lines"][-1]["id"]
    project = delete_line(db, last_id)
    assert all(ln["id"] != last_id for ln in project["lines"])


def test_http_pdf_import(tmp_path, monkeypatch):
    dbfile = tmp_path / "http.db"
    reset_db_singleton()
    monkeypatch.setenv("WEEKPLANNING_DB", str(dbfile))
    # Point default DB path
    import app.db as dbmod
    dbmod.DB_PATH = dbfile
    reset_db_singleton()
    client = TestClient(app)
    pdf = tmp_path / "a50.pdf"
    _write_pdf(pdf, _page("012510375A50", GOLD_A50))
    with pdf.open("rb") as f:
        res = client.post("/api/import/pdf", files={"file": ("a50.pdf", f, "application/pdf")})
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["number"] == "012510375A50"
    assert len(body["lines"]) == 20
    reset_db_singleton()
