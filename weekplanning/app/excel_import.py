"""Seed / import profile catalog from Excel.

Two layouts are supported:

1. Zaagmatrix (as described by the shop): row 1 machine names Quadra / BJM /
   Schirmer / Comet, row 2 Beperkt / Niet, column A profile numbers, ``x`` flags,
   Comet follow-up codes Q / B / S / QS / X.
2. Inrichting-sheet (the file currently supplied as Inrichting_Elace.xlsx):
   column A series codes (RT, AA4410, …), numbered process-step columns, cells
   with routing codes. Comet in a routing cell → comet_followup = X.
"""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from .db import ROOT, Database, get_db, utcnow

SEED_XLSX = ROOT / "seed" / "Inrichting_Elace.xlsx"
COMET_CODES = {"Q", "B", "S", "QS", "X"}
MACHINE_ALIASES = {
    "quadra": "quadra",
    "bjm": "bjm",
    "schirmer": "schirmer",
    "comet": "comet",
}


def seed_if_needed(xlsx_path: Path | None = None, db: Database | None = None) -> dict:
    path = xlsx_path or SEED_XLSX
    db = db or get_db()
    row = db.fetchone("SELECT value FROM meta WHERE key = 'excel_seeded'")
    if row and row["value"] == "1":
        return {"seeded": False, "reason": "already"}
    result = import_excel(path, db=db)
    db.execute(
        "INSERT INTO meta(key, value) VALUES ('excel_seeded', '1') "
        "ON CONFLICT(key) DO UPDATE SET value = '1'"
    )
    db.commit()
    result["seeded"] = True
    return result


def import_excel(xlsx_path: Path, db: Database | None = None) -> dict:
    own = db is None
    if own:
        db = get_db()
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    layout = detect_layout(ws)
    if layout == "zaagmatrix":
        result = _import_zaagmatrix(ws, db)
    else:
        result = _import_inrichting(ws, db)
    result["layout"] = layout
    result["file"] = str(xlsx_path)
    if own:
        db.commit()
    return result


def detect_layout(ws) -> str:
    row1 = [
        str(ws.cell(1, c).value or "").strip().lower()
        for c in range(1, min(ws.max_column, 20) + 1)
    ]
    joined = " ".join(row1)
    if "quadra" in joined or "schirmer" in joined or "bjm" in joined:
        return "zaagmatrix"
    return "inrichting"


def _is_x(value) -> bool:
    if value is None:
        return False
    return str(value).strip().lower() in {"x"}


def _import_zaagmatrix(ws, db: Database) -> dict:
    machines = {r["code"]: r["id"] for r in db.fetchall("SELECT id, code FROM machines")}
    colmap: dict[int, tuple[str, str]] = {}
    current_machine = None
    for c in range(2, ws.max_column + 1):
        header = str(ws.cell(1, c).value or "").strip()
        if header:
            token = header.lower().split()[0]
            key = MACHINE_ALIASES.get(token)
            if key:
                current_machine = key
            else:
                low = header.lower()
                for alias, code in MACHINE_ALIASES.items():
                    if alias in low:
                        current_machine = code
        sub = str(ws.cell(2, c).value or "").strip().lower()
        if current_machine == "comet":
            colmap[c] = ("comet", "comet")
        elif current_machine and sub.startswith("beperkt"):
            colmap[c] = (current_machine, "beperkt")
        elif current_machine and sub.startswith("niet"):
            colmap[c] = (current_machine, "niet")
        elif current_machine and not sub:
            colmap[c] = (current_machine, "niet")

    imported = 0
    start_row = 3 if any(str(ws.cell(2, c).value or "") for c in range(1, 12)) else 2
    for r in range(start_row, ws.max_row + 1):
        number = ws.cell(r, 1).value
        if number is None or str(number).strip() == "":
            continue
        number = str(number).strip()
        if number.lower() in {"profiel", "profielnummer", "artikel"}:
            continue
        flags: dict[str, dict] = {
            "quadra": {"beperkt": 0, "niet": 0},
            "bjm": {"beperkt": 0, "niet": 0},
            "schirmer": {"beperkt": 0, "niet": 0},
        }
        comet_code = None
        for c, (mcode, kind) in colmap.items():
            val = ws.cell(r, c).value
            if kind == "comet" or mcode == "comet":
                token = str(val).strip().upper() if val is not None else ""
                token = token.replace(" ", "")
                if token in COMET_CODES:
                    comet_code = token
                elif _is_x(val):
                    comet_code = "X"
            elif kind in flags.get(mcode, {}):
                if _is_x(val):
                    flags[mcode][kind] = 1
        _upsert_profile(db, number, notes="", comet=comet_code, source="excel")
        pid = db.fetchone("SELECT id FROM profiles WHERE number = ?", (number,))["id"]
        for mcode, f in flags.items():
            db.execute(
                """
                INSERT INTO profile_machine (profile_id, machine_id, beperkt, niet)
                VALUES (?, ?, ?, ?)
                ON CONFLICT(profile_id, machine_id) DO UPDATE SET
                    beperkt = excluded.beperkt,
                    niet = excluded.niet
                """,
                (pid, machines[mcode], f["beperkt"], f["niet"]),
            )
        imported += 1
    return {"imported": imported}


def _import_inrichting(ws, db: Database) -> dict:
    seen: dict[str, str] = {}
    imported = 0
    for r in range(1, ws.max_row + 1):
        raw = ws.cell(r, 1).value
        if raw is None:
            continue
        number = str(raw).strip()
        if not number or re.fullmatch(r"\d+", number):
            continue
        steps: list[str] = []
        has_comet = False
        for c in range(2, ws.max_column + 1):
            val = ws.cell(r, c).value
            if val is None or str(val).strip() == "":
                continue
            text = str(val).replace("\n", " / ").strip()
            steps.append(f"{c - 1}: {text}")
            if re.search(r"\bcomet\b", text, re.I):
                has_comet = True
        notes = "Inrichting: " + " · ".join(steps)
        if number in seen:
            if notes != seen[number]:
                db.execute(
                    "UPDATE profiles SET notes = notes || char(10) || ? WHERE number = ?",
                    ("Variant: " + notes, number),
                )
            continue
        comet = "X" if has_comet else None
        _upsert_profile(db, number, notes=notes, comet=comet, source="excel")
        seen[number] = notes
        imported += 1
    return {"imported": imported, "note": "inrichting-layout (geen zaagmatrix)"}


def _upsert_profile(db: Database, number: str, notes: str, comet: str | None, source: str) -> None:
    db.execute(
        """
        INSERT INTO profiles (number, notes, comet_followup, source, updated_at)
        VALUES (?, ?, ?, ?, ?)
        ON CONFLICT(number) DO UPDATE SET
            notes = CASE
                WHEN excluded.notes != '' THEN excluded.notes
                ELSE profiles.notes
            END,
            comet_followup = excluded.comet_followup,
            source = excluded.source,
            updated_at = excluded.updated_at
        """,
        (number, notes, comet, source, utcnow()),
    )
